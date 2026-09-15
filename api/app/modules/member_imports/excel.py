from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from app.modules.member_imports.schemas import MAX_IMPORT_SIZE


MAX_FILE_SIZE = 2 * 1024 * 1024
MAX_UNCOMPRESSED_FILE_SIZE = 10 * 1024 * 1024
EMAIL_HEADERS = {"email", "e-mail", "adresse email", "adresse e-mail"}


class InvalidMemberImportFileError(ValueError):
    pass


def extract_emails_from_excel(content: bytes) -> list[str]:
    if not content:
        raise InvalidMemberImportFileError("The Excel file is empty")
    if len(content) > MAX_FILE_SIZE:
        raise InvalidMemberImportFileError("The Excel file exceeds 2 MB")

    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_UNCOMPRESSED_FILE_SIZE:
                raise InvalidMemberImportFileError(
                    "The uncompressed Excel file exceeds 10 MB"
                )
    except BadZipFile as exc:
        raise InvalidMemberImportFileError("The Excel file is invalid") from exc

    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise InvalidMemberImportFileError("The Excel file is invalid") from exc

    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise InvalidMemberImportFileError("The Excel file is empty")

        email_column = next(
            (
                index
                for index, value in enumerate(header)
                if str(value or "").strip().lower() in EMAIL_HEADERS
            ),
            None,
        )
        if email_column is None:
            raise InvalidMemberImportFileError(
                "An 'email' column is required"
            )

        emails: list[str] = []
        for row in rows:
            value = row[email_column] if email_column < len(row) else None
            if value is None or not str(value).strip():
                continue
            emails.append(str(value).strip())
            if len(emails) > MAX_IMPORT_SIZE:
                raise InvalidMemberImportFileError(
                    f"An import cannot exceed {MAX_IMPORT_SIZE} accounts"
                )

        if not emails:
            raise InvalidMemberImportFileError("No email address was found")
        return emails
    finally:
        workbook.close()
